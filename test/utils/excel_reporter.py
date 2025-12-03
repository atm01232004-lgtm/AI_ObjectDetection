import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime


class ExcelReporter:
    def __init__(self, report_folder="Test_Reports", report_title="Test Report"):
        self.report_folder = report_folder
        self.report_title = report_title
        self.results = []  # Danh sách chứa kết quả (Time, Name, Desc, Status, Error)

    def add_result(self, test_name, description, status, error_msg=""):
        """Thêm một dòng kết quả vào bộ nhớ đệm"""
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.results.append([timestamp, test_name, description, status, error_msg])

    def save_report(self):
        """Xuất file Excel ra thư mục"""
        print(f"\n>>> Đang xuất báo cáo vào thư mục: {self.report_folder}...")

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Report"

            # 1. Header
            headers = ["Thời gian", "Tên Test Case", "Mô tả chức năng", "Trạng thái", "Chi tiết lỗi"]
            ws.append(headers)

            # Style Header
            header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            header_font = Font(bold=True)
            for col_num, _ in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            # 2. Ghi dữ liệu & Tô màu
            pass_font = Font(color="008000", bold=True)
            fail_font = Font(color="FF0000", bold=True)

            for row_data in self.results:
                ws.append(row_data)
                current_row = ws.max_row

                # Cột trạng thái là cột 4
                status_cell = ws.cell(row=current_row, column=4)
                if status_cell.value == "PASS":
                    status_cell.font = pass_font
                else:
                    status_cell.font = fail_font

            # 3. Chỉnh độ rộng cột
            widths = [20, 30, 40, 15, 50]
            for i, w in enumerate(widths, 1):
                col_letter = openpyxl.utils.get_column_letter(i)
                ws.column_dimensions[col_letter].width = w

            # 4. Lưu file
            if not os.path.exists(self.report_folder):
                os.makedirs(self.report_folder)

            filename = f"{self.report_title}_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.xlsx"
            full_path = os.path.join(self.report_folder, filename)

            wb.save(full_path)
            print(f">>> Đã lưu báo cáo tại: {os.path.abspath(full_path)}")

        except Exception as e:
            print(f"❌ Lỗi xuất Excel: {e}")